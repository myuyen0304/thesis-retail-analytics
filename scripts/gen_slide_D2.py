# -*- coding: utf-8 -*-
"""Sinh bo slide trinh bay bai toan D2 — Xoi mon nen khach hang.

Chay:  .venv/Scripts/python.exe scripts/gen_slide_D2.py
Ra:    D2-Slides.pptx  (23 slide, 16:9)

Moi con so trong file nay lay tu docs/tong-hop-D2.md va
docs/phan-bien/cau-hoi-phan-bien-can-kiem.md — khong go tay so moi.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Bang mau va chu ───────────────────────────────────────────────────────────
INK   = RGBColor(0x1A, 0x20, 0x2C)
MUTED = RGBColor(0x5A, 0x65, 0x70)
ACC   = RGBColor(0x2B, 0x6C, 0xB0)
DEEP  = RGBColor(0x1A, 0x36, 0x5D)
WARN  = RGBColor(0xB7, 0x53, 0x09)
BAD   = RGBColor(0xC5, 0x30, 0x30)
GOOD  = RGBColor(0x2F, 0x85, 0x5A)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
PANEL2= RGBColor(0xE8, 0xEF, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT  = 'Segoe UI'
MONO  = 'Consolas'

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width  = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

_n = [0]


# ── Ha tang ───────────────────────────────────────────────────────────────────
def noline(shp):
    shp.line.fill.background()
    return shp


def rect(s, x, y, w, h, fill, rounded=False, adj=0.05):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.shadow.inherit = False
    if rounded:
        shp.adjustments[0] = adj
    return noline(shp)


def _runs(p, text, size, color, bold, hl, font, italic):
    """Tach **dam** va `ma nguon` thanh cac run rieng."""
    for i, seg in enumerate(text.split('**')):
        if not seg:
            continue
        strong = (i % 2 == 1)
        for j, part in enumerate(seg.split('`')):
            if not part:
                continue
            code = (j % 2 == 1)
            r = p.add_run()
            r.text = part
            f = r.font
            f.name = MONO if code else font
            f.size = Pt(size * 0.92 if code else size)
            f.italic = italic
            f.bold = bold or strong
            if strong and hl is not None:
                f.color.rgb = hl
            elif code and color == INK:
                f.color.rgb = DEEP
            else:
                f.color.rgb = color


def tb(s, x, y, w, h, lines, size=17, color=INK, bold=False, hl=None,
       align=PP_ALIGN.LEFT, space=8, spacing=1.22, anchor=MSO_ANCHOR.TOP,
       font=FONT, italic=False, bullet=False):
    """lines: str hoac list[str]. Ho tro **dam** trong noi dung."""
    if isinstance(lines, str):
        lines = [lines]
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space)
        if bullet:
            _runs(p, '• ', size, ACC, True, None, font, False)
        _runs(p, ln, size, color, bold, hl, font, italic)
    return box


def head(s, kicker, title, sub=None):
    """Tieu de trang chuan + so trang."""
    _n[0] += 1
    rect(s, 0, 0, W, 0.11, ACC)
    if kicker:
        tb(s, 0.75, 0.52, 11.8, 0.3, kicker.upper(), size=11.5, color=ACC,
           bold=True, space=0)
    two = len(title) > 48          # tieu de dai -> hai dong, chu nho lai
    th = 1.04 if two else 0.62
    tb(s, 0.75, 0.86, 11.8, th, title, size=26 if two else 30, color=DEEP,
       bold=True, space=0, spacing=1.1)
    y = 0.86 + th + 0.08
    if sub:
        tb(s, 0.75, y, 11.8, 0.42, sub, size=14.5, color=MUTED, space=0)
        y += 0.44
    rect(s, 0.75, y + 0.04, 1.5, 0.035, ACC)
    tb(s, 12.05, 6.92, 0.55, 0.3, str(_n[0]), size=11, color=MUTED,
       align=PP_ALIGN.RIGHT, space=0)
    return y + 0.36


def stat(s, x, y, w, num, label, color=ACC, numsize=34, h=1.32, bg=PANEL):
    rect(s, x, y, w, h, bg, rounded=True, adj=0.08)
    tb(s, x + 0.18, y + 0.17, w - 0.36, 0.6, num, size=numsize, color=color,
       bold=True, align=PP_ALIGN.CENTER, space=0)
    tb(s, x + 0.18, y + h - 0.52, w - 0.36, 0.45, label, size=11.5, color=MUTED,
       align=PP_ALIGN.CENTER, space=0, spacing=1.1)


def stats(s, y, items, x0=0.75, total=11.83, gap=0.22, **kw):
    """items: (so, nhan[, mau[, co chu rieng]])."""
    n = len(items)
    w = (total - gap * (n - 1)) / n
    for i, it in enumerate(items):
        k = dict(kw)
        if len(it) > 3:
            k['numsize'] = it[3]
        stat(s, x0 + i * (w + gap), y, w, it[0], it[1],
             color=it[2] if len(it) > 2 else ACC, **k)


def table(s, x, y, w, headers, rows, colw=None, fs=12.5, hfs=11.5, rh=0.33,
          hrh=0.36, align=None, hl=None, rowcol=None):
    """Bang phang, khong vien, header nen xanh."""
    nr, nc = len(rows) + 1, len(headers)
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                            Inches(hrh + rh * len(rows)))
    t = gf.table
    # bo style mac dinh (khong ke, khong soc)
    tbl = gf._element.graphic.graphicData.tbl
    tbl[0][-1].text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    t.first_row = False
    t.horz_banding = False

    if colw:
        tot = sum(colw)
        for i, cw in enumerate(colw):
            t.columns[i].width = Inches(w * cw / tot)
    t.rows[0].height = Inches(hrh)
    for r in range(1, nr):
        t.rows[r].height = Inches(rh)

    align = align or [PP_ALIGN.LEFT] * nc

    def put(cell, text, size, color, bold, fill, al):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = al
        p.line_spacing = 1.08
        _runs(p, str(text), size, color, bold, hl, FONT, False)

    for c, htxt in enumerate(headers):
        put(t.cell(0, c), htxt, hfs, WHITE, True, ACC, align[c])
    for r, row in enumerate(rows):
        bg = WHITE if r % 2 == 0 else PANEL
        for c, val in enumerate(row):
            col = INK
            if rowcol and rowcol[r] is not None and c == len(row) - 1:
                col = rowcol[r]
            put(t.cell(r + 1, c), val, fs, col, False, bg, align[c])
    return gf


def mono(s, x, y, w, h, text, size=14, color=DEEP, bg=PANEL):
    rect(s, x, y, w, h, bg, rounded=True, adj=0.04)
    tb(s, x + 0.3, y + 0.22, w - 0.6, h - 0.44, text.split('\n'), size=size,
       color=color, font=MONO, space=2, spacing=1.16)


def callout(s, x, y, w, h, text, color=ACC, bg=PANEL2, size=14.5, icon=None):
    rect(s, x, y, w, h, bg, rounded=True, adj=0.06)
    rect(s, x, y, 0.07, h, color)
    tb(s, x + 0.32, y + 0.2, w - 0.6, h - 0.4, text.split('\n'), size=size,
       color=INK, hl=color, space=5, spacing=1.25, anchor=MSO_ANCHOR.MIDDLE)


def new():
    return prs.slides.add_slide(BLANK)


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · BÌA
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
rect(s, 0, 0, W, H, DEEP)
rect(s, 0, 0, 0.22, H, ACC)
tb(s, 1.25, 1.72, 10.8, 0.35, 'KHÓA LUẬN TỐT NGHIỆP  ·  NGÀNH HỆ THỐNG THÔNG TIN',
   size=12.5, color=RGBColor(0x8F, 0xB8, 0xE0), bold=True, space=0)
tb(s, 1.25, 2.24, 11, 1.0, 'Xói mòn nền khách hàng', size=46, color=WHITE,
   bold=True, space=0)
tb(s, 1.25, 3.30, 10.6, 0.75,
   'Bài toán D2 trong đề tài  ·  Dự báo doanh thu và giá vốn hàng bán theo ngày',
   size=18, color=RGBColor(0xC5, 0xDA, 0xEF), space=0, spacing=1.3)
rect(s, 1.25, 4.28, 1.6, 0.035, ACC)
tb(s, 1.25, 4.62, 11, 1.0,
   ['Dữ liệu thương mại điện tử thời trang  ·  04/07/2012 → 31/12/2022',
    '121.930 khách hàng  ·  646.945 đơn  ·  3.833 ngày'],
   size=14.5, color=RGBColor(0x9E, 0xBE, 0xDC), space=4, spacing=1.3)
tb(s, 1.25, 6.32, 11, 0.4, 'Trường Đại học Công nghiệp Thành phố Hồ Chí Minh',
   size=13, color=RGBColor(0x7B, 0xA3, 0xCC), space=0)
_n[0] = 1

# ═══════════════════════════════════════════════════════════════════════════════
# 2 · NỘI DUNG
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Nội dung trình bày', 'Bốn phần',
         'Từ bài toán → kết quả số → bộ chỉ số → tự phản biện')
parts = [
    ('I',   'Bài toán và cách phân rã',
     'Doanh thu mất 44,4% từ đỉnh. Phân rã ba tầng để tìm chỗ hỏng.', 'Slide 3–5'),
    ('II',  'Kết quả demo — sáu bài toán nhỏ',
     'Số liệu chạy trực tiếp trên dữ liệu gốc: BTN1 → BTN4, chín giả thuyết.', 'Slide 6–12'),
    ('III', 'Bộ chỉ số và kiểm chứng',
     '13 Measure → 13 Metric → 7 KPI. Sáu phép kiểm chéo.', 'Slide 13–14'),
    ('IV',  'Tự phản biện 31 mục',
     'Tám lỗi tìm được trong chính tài liệu của mình, và cách sửa.', 'Slide 15–23'),
]
yy = y + 0.05
for num, title, desc, rng in parts:
    rect(s, 0.75, yy, 11.83, 0.96, PANEL, rounded=True, adj=0.08)
    rect(s, 0.75, yy, 0.07, 0.96, ACC)
    tb(s, 1.05, yy + 0.18, 0.8, 0.6, num, size=24, color=ACC, bold=True, space=0)
    tb(s, 1.95, yy + 0.15, 7.9, 0.34, title, size=16.5, color=DEEP, bold=True, space=0)
    tb(s, 1.95, yy + 0.52, 8.4, 0.38, desc, size=12.5, color=MUTED, space=0)
    tb(s, 10.6, yy + 0.32, 1.8, 0.35, rng, size=12, color=ACC, bold=True,
       align=PP_ALIGN.RIGHT, space=0)
    yy += 1.08

# ═══════════════════════════════════════════════════════════════════════════════
# 3 · BÀI TOÁN LỚN
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần I · Bài toán', 'Doanh thu giảm — nhưng giảm ở đâu?',
         'Đỉnh 2016 → 2022, bộ lọc ALL, tất cả cùng một mốc')
stats(s, y + 0.02, [
    ('−44,4%', 'DOANH THU', BAD),
    ('−56,2%', 'SỐ ĐƠN — đây mới là chỗ mất', BAD),
    ('+27,0%', 'AOV — giá trị mỗi đơn LẠI TĂNG', GOOD),
])
tb(s, 0.75, y + 1.56, 11.83, 1.0, [
    'Mức giảm **không** đến từ giá trị mỗi đơn — AOV còn tăng. Nó đến từ **số đơn**.',
    'Nhưng số đơn là **kết quả**, không phải nguyên nhân: đơn hàng do khách hàng tạo ra. '
    'Dừng ở đây thì chỉ mô tả được triệu chứng.',
], size=15, hl=DEEP, space=8)
callout(s, 0.75, y + 2.58, 11.83, 0.78,
        'Câu hỏi thật: nền khách hàng xói mòn ở **khâu nào**, **cơ chế gì** gây ra, và '
        '**can thiệp nào** giữ lại được?', color=ACC, size=15)
callout(s, 0.75, y + 3.46, 11.83, 0.98,
        '⚠  Một chỗ lệch nữa, tìm ra khi dựng slide (ngoài 31 mục đã kiểm): tài liệu ghép '
        '**−44,4% (mốc 2016)** với **AOV +50,7% (mốc 2013)** trong cùng một câu.\n'
        'Cùng mốc 2016 thì AOV chỉ tăng **27,0%**. Phải nói rõ mốc, hoặc đưa cả ba về một mốc.',
        color=WARN, bg=RGBColor(0xFD, 0xF3, 0xE7), size=13)

# ═══════════════════════════════════════════════════════════════════════════════
# 4 · TRỤC PHÂN RÃ BA TẦNG
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần I · Phương pháp', 'Trục phân rã ba tầng',
         'Sáu bài toán nhỏ ứng đúng các thành phần trong phân rã')
mono(s, 0.75, y + 0.02, 11.83, 1.72,
     'Doanh thu\n'
     '  = Số đơn                       ×  AOV\n'
     '  = (Khách hoạt động × Tần suất)  ×  AOV\n'
     '  = ((Khách mới + Khách giữ lại) × Tần suất)  ×  AOV', size=15)
tb(s, 0.75, y + 1.94, 11.83, 0.5,
   'Vì các bài toán nhỏ bám vào **thành phần của phép nhân** chứ không gom theo chủ đề, '
   'chúng chia hết không gian nguyên nhân theo **định nghĩa**.', size=15, hl=DEEP, space=0)
callout(s, 0.75, y + 2.64, 11.83, 1.46,
        '⚠  Tự phản biện A9 — nói "phủ kín nguyên nhân" là **quá lời**.\n'
        'Đẳng thức đúng vì định nghĩa, nhưng **AOV không bài toán nhỏ nào phụ trách**, và nhánh '
        '**"giành lại" = 52,9%** khách hoạt động 2022 bị gộp chìm vào "khách giữ lại". '
        'Phải nói: phủ kín **theo số học**, chưa phủ kín **theo trách nhiệm phân tích**.',
        color=WARN, bg=RGBColor(0xFD, 0xF3, 0xE7), size=14)

# ═══════════════════════════════════════════════════════════════════════════════
# 5 · CÂY BÀI TOÁN
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần I · Cấu trúc', 'Sáu bài toán nhỏ và tiến độ',
         'Bốn bài toán đã có kết quả · hai bài còn treo, và treo vì lý do khác nhau')
table(s, 0.75, y + 0.02, 11.83,
      ['', 'Bài toán nhỏ', 'Kết quả', 'Trạng thái'],
      [['BTN1', 'Tập khách đang ở trạng thái nào?',
        '31.684 chưa mua · 65.493 ngủ đông · 24.753 hoạt động', '✔ Xong'],
       ['BTN2', 'Mất khách hay mua thưa đi?',
        'Ít khách 72,2% · mua thưa 45,2% · tương tác −17,4%', '✔ Xong'],
       ['BTN3', 'Kích hoạt hỏng hay giữ chân hỏng?',
        'Rổ cạn 36,4% · tỷ lệ hút giảm 63,6%  (đã sửa)', '✔ Xong'],
       ['BTN4', 'Cơ chế ở đơn đầu là gì?',
        'Khuyến mãi: thô 39,2% → kiểm soát xong còn 5,2%', '★ Trọng tâm'],
       ['BTN5', 'Nền khách đỡ nổi 2023–24?',
        'Đi ngang ba năm, nhưng n = 3 quá mỏng', '◐ Cần Chow test'],
       ['BTN6', 'Ngân sách nên đi đâu?',
        'Đã loại phương án theo kênh; thiếu bảng chi phí', '✖ Thiếu dữ liệu']],
      colw=[0.9, 3.4, 5.6, 1.9], fs=12.5, rh=0.44,
      align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER],
      rowcol=[GOOD, GOOD, GOOD, ACC, WARN, BAD])
callout(s, 0.75, y + 3.16, 11.83, 0.95,
        'BTN5 treo vì **chưa làm** — công cụ có sẵn, làm được ngay.   '
        'BTN6 treo vì **bộ dữ liệu không có bảng chi phí marketing**.\n'
        'Phân biệt này quan trọng: một cái là thiếu công, một cái là **giới hạn dữ liệu**.',
        color=ACC, size=14)

# ═══════════════════════════════════════════════════════════════════════════════
# 6 · BTN1
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · BTN1', 'Tập khách đang ở trạng thái nào?',
         'Phân ba trạng thái trên danh sách đóng 121.930 khách')
stats(s, y + 0.02, [
    ('31.684', 'CHƯA TỪNG MUA\nđăng ký rồi im lặng', MUTED),
    ('65.493', 'NGỦ ĐÔNG\ntừng mua, đã ngưng', WARN),
    ('24.753', 'HOẠT ĐỘNG\ncòn mua trong 2022', GOOD),
], h=1.45)
tb(s, 0.75, y + 1.66, 11.83, 0.44,
   'Phép kiểm tổng:   31.684 + 65.493 + 24.753 = **121.930**  — đúng bằng số dòng '
   '`customers.csv`. Không ai rơi ra ngoài, không ai bị đếm hai lần.',
   size=14.5, hl=DEEP, space=0)
callout(s, 0.75, y + 2.26, 11.83, 1.84,
        '⚠  Tự phản biện B15 — hai chỉ số dùng **hai bộ lọc khác nhau** mà không ghi nhãn.\n'
        'Me1 ("chưa từng mua") tính trên `live` — loại đơn `cancelled`. '
        'Me8a tính trên `ALL` — giữ đơn `cancelled`.\n'
        'Hệ quả: cộng lại ra **98,26%** thay vì 100%. Khách chỉ có đơn `cancelled` **rơi mất**.\n'
        'Con số nhất quán là **27,73%**, không phải 26,0%.',
        color=BAD, bg=RGBColor(0xFD, 0xEC, 0xEC), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 7 · BTN2
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · BTN2', 'Mất khách, hay người ở lại mua thưa đi?',
         'Phân rã ba thành phần: số đơn = khách hoạt động × tần suất')
stats(s, y + 0.02, [
    ('72,2%', 'do ÍT KHÁCH hơn', BAD),
    ('45,2%', 'do MUA THƯA hơn', WARN),
    ('−17,4%', 'số hạng tương tác', MUTED),
])
tb(s, 0.75, y + 1.6, 11.83, 0.95, [
    'Hai sự cố xảy ra **đồng thời** và **nhân lên nhau** — nên tổng ba phần vượt 100%, '
    'phần dôi bị số hạng tương tác kéo lại.',
    '**Hệ quả thực hành:** sửa một trong hai không đủ. Giữ chân tốt hơn mà nền khách vẫn co '
    'thì doanh thu vẫn giảm.',
], size=15, hl=DEEP, space=8)
callout(s, 0.75, y + 2.72, 11.83, 1.38,
        '⚠  Tự phản biện B6 — **hai tầng dùng hai phương pháp phân rã khác nhau.**\n'
        'Tầng này dùng phân rã **số học** (có số hạng tương tác), tầng BTN3 dùng phân rã '
        '**logarit** (không có). Tính lại BTN2 bằng logarit cho **63,8% / 36,2%**.\n'
        'Thứ tự **không đổi** — H1 vẫn đứng — nhưng tài liệu **phải khai báo** dùng phương pháp nào.',
        color=WARN, bg=RGBColor(0xFD, 0xF3, 0xE7), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 8 · BTN3
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · BTN3', 'Kích hoạt hỏng, hay chỉ là rổ khách đã cạn?',
         'Phân rã logarit  ln(Khách mới) = ln(Pool) + ln(Tỷ lệ hút)')
tb(s, 0.75, y + 0.02, 11.83, 0.5,
   '`customers.csv` là **danh sách đóng** 121.930 người. Rổ chưa mua chỉ co lại theo thời gian — '
   'nên "khách mới giảm" **trộn lẫn** hiệu ứng cơ học với tín hiệu thật. Phân rã logarit tách đôi.',
   size=14.5, hl=DEEP, space=0)
stats(s, y + 0.72, [
    ('36,4%', 'do RỔ CẠN\nhiệu ứng cơ học', MUTED),
    ('63,6%', 'do TỶ LỆ HÚT SỤP\ntín hiệu thật', BAD),
    ('24,09% → 3,78%', 'Tỷ lệ hút từ pool\n2013 → 2022', BAD, 25),
], h=1.45)
_last = s.shapes[-1]
callout(s, 0.75, y + 2.4, 11.83, 1.72,
        '✖  Tự phản biện B1 — **tài liệu ghi 26,9% / 73,1%, con số đó SAI.**\n'
        'Vòng lặp tính Pool gán giá trị **trước khi trừ** và bắt đầu từ năm 2013, làm '
        '**rơi mất toàn bộ cohort 2012**. Tính lại đúng: **36,4% / 63,6%**.\n'
        '**Kết luận không đổi** — tỷ lệ hút vẫn là nguyên nhân chính — nhưng rổ cạn '
        'nặng hơn tài liệu thừa nhận gần **10 điểm phần trăm**.',
        color=BAD, bg=RGBColor(0xFD, 0xEC, 0xEC), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 9 · BTN4 — COX
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · BTN4 ★ Trọng tâm', 'Cơ chế ở đơn hàng đầu tiên',
         'Mô hình Cox proportional hazards · 84.566 khách · 74,1% có sự kiện · '
         'trung vị 312 ngày')
table(s, 0.75, y + 0.02, 11.83,
      ['Biến ở đơn đầu', 'HR', 'p', 'Đọc thế nào'],
      [['`cohort_year` — năm gia nhập', '0,7478', '< 0,001',
        'Mỗi năm gia nhập muộn hơn → quay lại chậm hơn 25%'],
       ['`promo_first` — có khuyến mãi', '0,948', '< 0,001',
        'Chậm hơn ~5,2% sau khi kiểm soát — nhỏ'],
       ['`log_aov` — giá trị đơn đầu', '—', '< 0,001',
        'Có ý nghĩa nhưng VI PHẠM giả định Cox'],
       ['`delivery_days` — giao chậm', '≈ 1,00', '0,303',
        'Không có tín hiệu'],
       ['`returned_first` — trả hàng', '≈ 1,00', '0,161',
        'Không có tín hiệu']],
      colw=[3.3, 1.2, 1.1, 6.23], fs=12.5, rh=0.42,
      align=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.LEFT])
callout(s, 0.75, y + 2.52, 11.83, 1.58,
        '⚠  Tự phản biện A8 — **concordance = 0,653**.\n'
        'Đoán mò hoàn toàn cho 0,50. Mô hình chỉ hơn đoán mò **15 điểm** — '
        'đủ để nói về **cơ chế và thứ tự quan trọng**, nhưng **không nên gọi là "mô hình dự đoán"**.\n'
        'Schoenfeld: `log_aov` và `category_Outdoor` vi phạm giả định. Hai biến mang kết luận chính '
        '(`promo_first` p = 0,104 · `cohort_year` p = 0,269) **không vi phạm**.',
        color=WARN, bg=RGBColor(0xFD, 0xF3, 0xE7), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 10 · KAPLAN–MEIER
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · BTN4', 'Đường sống sót Kaplan–Meier',
         'Xác suất CHƯA quay lại mua, theo thời gian kể từ đơn đầu')
img = os.path.join('docs', 'hinh', '12-kaplan-meier-promo.png')
if os.path.exists(img):
    s.shapes.add_picture(img, Inches(0.6), Inches(y + 0.02), height=Inches(3.9))
else:
    mono(s, 0.6, y + 0.02, 6.5, 3.9, '[ docs/hinh/12-kaplan-meier-promo.png ]')
tb(s, 7.5, y + 0.1, 5.08, 0.3, 'ĐỌC BIỂU ĐỒ NÀY THẾ NÀO', size=11.5, color=ACC,
   bold=True, space=0)
tb(s, 7.5, y + 0.52, 5.08, 2.4, [
    'Trục dọc: tỷ lệ khách **chưa** quay lại. Đường **thấp hơn** = quay lại **nhiều hơn**.',
    'Tại mốc 1 năm, hai nhóm chênh **10,0 điểm phần trăm**.',
    'Log-rank: χ² = 873,8 · p < 0,001 — khác biệt **có ý nghĩa thống kê**.',
], size=13.5, hl=DEEP, space=9, bullet=True)
callout(s, 7.5, y + 2.86, 5.08, 1.06,
        'Nhưng khoảng cách thô này **gồm cả hiệu ứng cohort**. Sau khi kiểm soát, '
        'HR chỉ còn **0,948** — xem slide 19.', color=WARN,
        bg=RGBColor(0xFD, 0xF3, 0xE7), size=13)

# ═══════════════════════════════════════════════════════════════════════════════
# 11 · PHÁT HIỆN ⭐
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · Phát hiện chính', '★  Không phải khuyến mãi — là hiệu ứng cohort',
         'Đóng góp riêng của khóa luận: cùng một dữ liệu, hai kết luận ngược nhau')
bw = 5.6
rect(s, 0.75, y + 0.02, bw, 1.62, RGBColor(0xFD, 0xEC, 0xEC), rounded=True, adj=0.07)
tb(s, 1.0, y + 0.18, bw - 0.5, 0.3, 'TÍN HIỆU THÔ — chỉ mô tả', size=11.5,
   color=BAD, bold=True, space=0)
tb(s, 1.0, y + 0.52, bw - 0.5, 0.5, '−39,2%', size=34, color=BAD, bold=True, space=0)
tb(s, 1.0, y + 1.12, bw - 0.5, 0.42,
   'Mua ít hơn: 4,59 vs 7,56 đơn trọn đời', size=12, color=MUTED, space=0)
tb(s, 6.5, y + 0.62, 0.6, 0.5, '→', size=30, color=MUTED, bold=True,
   align=PP_ALIGN.CENTER, space=0)
rect(s, 7.23, y + 0.02, bw, 1.62, RGBColor(0xE9, 0xF5, 0xEE), rounded=True, adj=0.07)
tb(s, 7.48, y + 0.18, bw - 0.5, 0.3, 'SAU KHI KIỂM SOÁT — cơ chế', size=11.5,
   color=GOOD, bold=True, space=0)
tb(s, 7.48, y + 0.52, bw - 0.5, 0.5, '−5,2%', size=34, color=GOOD, bold=True, space=0)
tb(s, 7.48, y + 1.12, bw - 0.5, 0.42,
   'Kiểm soát cohort, danh mục, giá trị đơn → HR = 0,948', size=12, color=MUTED, space=0)
tb(s, 0.75, y + 1.86, 11.83, 0.5,
   'Biến `cohort_year` có **HR = 0,7478** — mạnh hơn hẳn mọi biến khác. '
   'Thứ tưởng là tác hại của khuyến mãi, phần lớn là **khách gia nhập muộn vốn đã kém hơn**.',
   size=15, hl=DEEP, space=0)
callout(s, 0.75, y + 2.52, 11.83, 1.58,
        'Vì sao điều này quan trọng:\n'
        'Nếu dừng ở phân tích mô tả, kết luận sẽ là **"cắt ngân sách khuyến mãi ngay"** — '
        'một quyết định **sai**, xuất phát từ nhầm tương quan với nhân quả.\n'
        'Bước đi từ **mô tả** sang **cơ chế** không phải làm cho đẹp bài. Nó **đổi hẳn khuyến nghị**.',
        color=ACC, size=14.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 12 · CHÍN GIẢ THUYẾT
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần II · Tổng kết', 'Chín giả thuyết — ghi cả cái bị bác bỏ',
         'Giả thuyết bị bác bỏ là KẾT QUẢ, không phải thất bại')
table(s, 0.75, y + 0.02, 11.83,
      ['#', 'Giả thuyết', 'Kết quả', 'Sau phản biện'],
      [['H1', 'Mất khách lớn hơn giảm tần suất', '✔ 72,2% vs 45,2%', 'Giữ'],
       ['H2', 'Sau khi trừ rổ cạn, tỷ lệ hút vẫn là chính', '✔ 63,6% vs 36,4%', 'Sửa số — B1'],
       ['H3', 'Chất lượng cohort giảm đơn điệu', '✔ Giảm liên tục', 'Giữ'],
       ['H4', 'Gãy 2019 là hệ quả trễ của suy giảm cohort', '✖ BỊ BÁC BỎ', 'Đổi: treo → bác bỏ — B9'],
       ['H5', 'Kênh thu nạp không phân hóa giá trị', '⚠ F = 0,823 · p = 0,533', 'Giữ — đã thận trọng'],
       ['H6', '2020–22 chưa thấy dấu hiệu tiếp tục rơi', '⚠ Ủng hộ yếu, n = 3', 'Củng cố — B4'],
       ['H7', 'Khuyến mãi đơn đầu giảm khả năng quay lại', '✔ Đúng nhưng nhỏ — HR 0,948', 'Thu hẹp — B12, B17'],
       ['H8', 'Giao hàng chậm giảm khả năng quay lại', '⚠ p = 0,303 — rỗng', 'Đổi ✖ → ⚠ — B21'],
       ['H9', 'Trả hàng là tín hiệu rời bỏ mạnh nhất', '⚠ p = 0,161 — rỗng', 'Đổi ✖ → ⚠ — B21']],
      colw=[0.65, 4.5, 3.55, 3.13], fs=12, rh=0.335,
      align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
      rowcol=[GOOD, WARN, GOOD, BAD, GOOD, ACC, WARN, WARN, WARN])
callout(s, 0.75, y + 3.62, 11.83, 0.86,
        'H8 và H9 **không còn** đọc là "giao hàng không ảnh hưởng giữ chân" — xem slide 17: '
        'hai trường dữ liệu đó có dấu hiệu **gán ngẫu nhiên**.',
        color=WARN, bg=RGBColor(0xFD, 0xF3, 0xE7), size=14)

# ═══════════════════════════════════════════════════════════════════════════════
# 13 · BỘ CHỈ SỐ
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần III · Bộ chỉ số', '13 Measure → 13 Metric → 7 KPI',
         'Mỗi bài toán nhỏ truy được xuống một KPI có ngưỡng và có hành động')
table(s, 0.75, y + 0.02, 11.83,
      ['', 'Vấn đề phát hiện được', 'Đã sửa thành'],
      [['Me8', 'Lỗi grain — tử số trộn khách chưa mua (không có recency) với khách ngủ đông',
        'Tách đôi: Me8a 26,0% · Me8b 72,6%'],
       ['K2', 'Mục tiêu tăng trưởng đặt trên một rổ CHỈ CÓ THỂ CO LẠI — bất khả thi về cấu trúc',
        'Đổi sang tỷ lệ hút từ pool'],
       ['K7', 'Sáu KPI cũ đều một hướng — không cái nào chặn chi phí',
        'Thêm guardrail: đơn đầu có khuyến mãi ≤ 30%']],
      colw=[0.95, 6.5, 4.38], fs=12.5, rh=0.55,
      align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
callout(s, 0.75, y + 2.08, 11.83, 2.02,
        '✖  Tự phản biện B8 và B19 — bộ KPI vẫn còn ba chỗ hỏng:\n'
        '**K7 vi phạm ngưỡng của chính nó ngay ngày ban hành** — thực tế đang ở **30,57%**, '
        'trong khi ngưỡng đặt là ≤ 30%.\n'
        '**K1 đòi tăng 2,86 lần**, mốc 20% cách hiện tại **sáu năm** — không phải mục tiêu điều hành được.\n'
        '**K5 chỉ đo được sau 36 tháng** — quá trễ để lái. Đã đề xuất hai chỉ số **dẫn báo** thay thế.',
        color=BAD, bg=RGBColor(0xFD, 0xEC, 0xEC), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 14 · KIỂM CHỨNG
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần III · Kiểm chứng', 'Số liệu tự kiểm nhau bằng đường khác',
         'Chạy lại cùng một script không chứng minh gì — chỉ chứng minh code chạy hai lần ra cùng kết quả')
table(s, 0.75, y + 0.02, 11.83,
      ['Phép kiểm', 'Kết quả', 'Có thể vỡ không?'],
      [['Ba nhóm khách cộng lại bằng tổng đăng ký', '31.684 + 65.493 + 24.753 = 121.930',
        'Hằng đúng'],
       ['Phân rã logarit: rổ cạn + tỷ lệ hút = tổng', 'Sai số 4,4 × 10⁻¹⁶', 'Hằng đúng'],
       ['Hai đường phân rã độc lập gặp nhau', 'Phễu và vòng đời cùng ra −0,348', 'Hằng đúng'],
       ['`sales.csv` vs Σ(quantity × unit_price)', 'Tỷ lệ 1,000000 trên 3.833 ngày',
        '★ CÓ — kiểm thật'],
       ['`payment_value` = gross − discount', 'Khớp 100% trên 646.945 đơn', '★ CÓ — kiểm thật'],
       ['Nghiệm thu bộ chỉ số', '8/8 dòng OK', '—']],
      colw=[4.5, 4.6, 2.73], fs=12.5, rh=0.4,
      align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER],
      rowcol=[MUTED, MUTED, MUTED, GOOD, GOOD, MUTED])
callout(s, 0.75, y + 2.86, 11.83, 1.24,
        '✖  Tự phản biện B10 — **ba trong năm phép kiểm là hằng đúng**, không thể vỡ dù số liệu sai.\n'
        'Nặng nhất: phép kiểm tài liệu gọi là **"mạnh nhất"** (hai đường phân rã gặp nhau) thực ra '
        '**yếu nhất** — hai đường dùng chung dữ liệu đầu vào nên buộc phải gặp. '
        'Đã đề xuất phép kiểm thay thế **có thể vỡ**.',
        color=BAD, bg=RGBColor(0xFD, 0xEC, 0xEC), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 15 · PHẢN BIỆN — TỔNG QUAN
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · Tự phản biện', 'Tự kiểm 31 mục trên chính tài liệu của mình',
         'Mỗi mục một script riêng, chạy trực tiếp trên dữ liệu gốc — không tin lại kết quả cũ')
stats(s, y + 0.02, [
    ('13', 'XÁC NHẬN ĐÚNG\ntài liệu đứng vững', GOOD),
    ('10', 'PHẢI THU HẸP\nđúng nhưng nói quá', WARN),
    ('8', 'LỆCH — PHẢI SỬA\ncon số hoặc kết luận sai', BAD),
], h=1.42)
tb(s, 0.75, y + 1.64, 11.83, 0.95, [
    'Nguyên tắc đã áp dụng: mỗi con số phải kèm **bảng nguồn · cột · bộ lọc · grain**. '
    'Khi lệch so với tài liệu thì ghi thẳng **LỆCH**, không làm tròn cho khớp.',
    'Khi không xác định được thì viết **"chưa xác định được"** kèm lý do — **không đoán**.',
], size=14.5, hl=DEEP, space=7)
callout(s, 0.75, y + 2.76, 11.83, 1.34,
        'Vì sao đưa phần này vào bài báo cáo:\n'
        'Tám lỗi trên **đều do chính nhóm tìm ra**, trước khi hội đồng hỏi. '
        'Một tài liệu **biết chỗ yếu của mình ở đâu** đáng tin hơn một tài liệu **không có chỗ yếu nào**.',
        color=ACC, size=14.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 16 · TÁM LỖI
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · Kết quả', 'Tám chỗ lệch đã tìm ra và cách sửa',
         'Mỗi dòng là một con số hoặc một kết luận trong tài liệu KHÔNG đứng được khi kiểm lại')
table(s, 0.75, y + 0.02, 11.83,
      ['#', 'Chỗ trong tài liệu', 'Lệch cái gì'],
      [['B1', 'Mục 2.2 · rổ chưa mua M13',
        'Vòng lặp bỏ mất cohort 2012 → rổ cạn 26,9% phải thành **36,4%**'],
       ['B2', 'Mục 12.2 · Kiểm 2',
        'Hằng đúng — Pool triệt tiêu; cặp số bịa 7/3 vẫn khớp'],
       ['B7', 'Mục 9.3 · `cohort_year` tuyến tính',
        'BÁC BỎ — LR = 91,47 · p = 8,2e-16; hỏng từ 2020, lệch tới 2,40 lần'],
       ['B8', 'Mục 8 · ngưỡng K1, K7',
        'K7 = **30,57%** vi phạm ngưỡng ≤ 30% ngay ngày ban hành'],
       ['B9', 'Mục 5 ↔ 11 · giả thuyết H4',
        'H4 BỊ BÁC BỎ — mô phỏng **+0,68%** trong khi thực tế **−38,60%**'],
       ['B10', 'Mục 12.2 · cả 5 phép kiểm',
        '3/5 hằng đúng; phép kiểm gọi là "mạnh nhất" thực ra yếu nhất'],
       ['B15', 'Mục 3 · "chưa từng mua"',
        'Me1 dùng `live`, Me8a dùng `ALL` → cộng lại 98,26%. Đúng là **27,73%**'],
       ['B20', 'Mục 3, 5.1 · `signup_date`',
        'KHÔNG sửa được — Spearman 0,0023 (p = 0,49), SD 3,5 năm']],
      colw=[0.72, 3.5, 7.61], fs=12, rh=0.385, hl=BAD,
      align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
tb(s, 0.75, y + 3.68, 11.83, 0.42,
   'Cả tám đều đã có **đề xuất câu chữ cụ thể** để sửa, ghi trong '
   '`docs/phan-bien/cau-hoi-phan-bien-can-kiem.md`.', size=13.5, hl=DEEP, space=0)

# ═══════════════════════════════════════════════════════════════════════════════
# 17 · B21
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · Phát hiện nặng nhất · B21',
         'Bốn trường có dấu hiệu gán ngẫu nhiên',
         'Kiểm mỗi trường bằng một tương quan lẽ ra PHẢI tồn tại trong dữ liệu thật')
table(s, 0.75, y + 0.02, 11.83,
      ['Trường', 'Phép kiểm', 'Kết quả', 'Nghi ngẫu nhiên'],
      [['`reviews` ⟷ `returns`', 'Số đơn có CẢ HAI', '**0** trên kỳ vọng **6.208**', 'Rất mạnh'],
       ['`delivery_days`', 'ANOVA theo vùng', '4,499 / 4,499 / 4,500 · p = 0,985', 'Mạnh'],
       ['`signup_date`', 'Spearman với ngày mua đầu', 'ρ = 0,0023 · p = 0,487', 'Mạnh'],
       ['`acquisition_channel`', 'ANOVA LTV theo kênh', 'F = 0,823 · p = 0,533', 'Mạnh'],
       ['`rating`', 'Spearman với số đơn trọn đời', 'ρ = −0,0734 · p < 0,001', 'Yếu — có liên hệ']],
      colw=[2.9, 3.2, 3.85, 1.88], fs=12, rh=0.38, hl=BAD,
      align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER],
      rowcol=[BAD, BAD, BAD, BAD, MUTED])
callout(s, 0.75, y + 2.44, 11.83, 1.66,
        '111.369 đơn có review.  36.062 đơn có trả hàng.  Số đơn có **cả hai: 0**.\n'
        'Nếu hai việc độc lập nhau thì kỳ vọng là **6.208 đơn**. Quan sát 0 — xác suất ngẫu nhiên '
        '**thực tế bằng không**. Bộ sinh gán review và trả hàng vào **hai tập rời nhau**.\n'
        '**Hệ quả:** không thể kiểm "đơn bị trả có rating thấp hơn không" — một quan hệ hiển nhiên '
        'phải có trong dữ liệu thật. Đây không phải kết quả rỗng, đây là **dữ liệu không tồn tại**.',
        color=BAD, bg=RGBColor(0xFD, 0xEC, 0xEC), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 18 · B11
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · B11', 'Hiệu ứng cohort, hay hiệu ứng thời kỳ?',
         'HR = 0,7478 là kết luận trọng tâm của BTN4 — câu hỏi này đánh thẳng vào đó')
tb(s, 0.75, y + 0.02, 11.83, 0.52,
   '**Nếu là cohort:** người gia nhập muộn vốn đã kém hơn → phải sửa khâu **thu nạp**.\n'
   '**Nếu là thời kỳ:** một cú sốc năm 2019 làm **mọi người** kém đi → phải tìm **cú sốc đó**.',
   size=14.5, hl=DEEP, space=4)
stats(s, y + 0.88, [
    ('61,2% → 46,0%', 'Cohort 2013 — CÙNG MỘT NHÓM NGƯỜI\nrơi đúng vào năm 2019', BAD),
    ('0,74 → 1,16', 'HR của `cohort_year` LẬT NGƯỢC\nkhi thêm biến thời kỳ vào mô hình', BAD),
], h=1.5)
callout(s, 0.75, y + 2.52, 11.83, 1.6,
        'Kết luận: **CẢ HAI**, và tài liệu chỉ nói một nửa.\n'
        'Cohort 2013 là nhóm cố định — không ai vào, không ai ra. Họ rơi 15 điểm **đúng năm 2019**. '
        'Chuyện xảy ra với người đã ở sẵn trong hệ thống **không thể** là hiệu ứng cohort.\n'
        'Đây là bài toán **Age–Period–Cohort** kinh điển: ba trục cộng tuyến hoàn hảo, '
        'không tách được nếu không có ràng buộc thêm. **Phải khai báo giới hạn này.**',
        color=WARN, bg=RGBColor(0xFD, 0xF3, 0xE7), size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 19 · B17
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · B17', 'Có tín hiệu — nhưng lớn bằng nào?',
         'Quy hazard ratio ra số người thật, thứ hội đồng và doanh nghiệp thực sự cần biết')
stats(s, y + 0.02, [
    ('1,4', 'ĐIỂM PHẦN TRĂM\nhiệu ứng thật của khuyến mãi', WARN),
    ('367', 'NGƯỜI\ntrên tổng 26.759 khách có khuyến mãi', WARN),
    ('25,8×', 'LẦN\n`cohort_year` mạnh hơn `promo_first`', BAD),
], h=1.5)
tb(s, 0.75, y + 1.72, 11.83, 0.98, [
    '`promo_first` có p < 0,001 — **chắc chắn có tín hiệu**. Nhưng cỡ hiệu ứng chỉ **1,4 điểm phần trăm**, '
    'tức là **367 người** trong số 26.759.',
    '`cohort_year` đổi **35,4 điểm phần trăm** — mạnh hơn **25,8 lần**.',
], size=14.5, hl=DEEP, space=7)
callout(s, 0.75, y + 2.72, 11.83, 1.38,
        'Vì sao phải trình bày cỡ hiệu ứng chứ không chỉ p-value:\n'
        'Với **84.566 quan sát**, gần như **mọi** biến đều sẽ có p < 0,001. p-value chỉ trả lời '
        '"có khác 0 không", **không** trả lời "có đáng để làm gì không".\n'
        'Một can thiệp đổi được **367 người** không xứng với một chương trình khuyến mãi toàn hệ thống.',
        color=ACC, size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 20 · KẾT LUẬN NÀO ĐỔI
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · Tổng kết', 'Sau phản biện, kết luận nào còn đứng?',
         'Chia theo cái mà kết luận DỰA VÀO — cấu trúc thời gian, hay một trường có thể ngẫu nhiên')
tb(s, 0.75, y + 0.02, 5.75, 0.32, 'CÒN ĐỨNG VỮNG', size=12, color=GOOD, bold=True, space=0)
rect(s, 0.75, y + 0.38, 5.75, 2.72, RGBColor(0xE9, 0xF5, 0xEE), rounded=True, adj=0.04)
tb(s, 1.0, y + 0.56, 5.25, 2.4, [
    'Phân rã số đơn = khách × tần suất (BTN2, H1)',
    'Pool và tỷ lệ hút (BTN3, H2) — sau khi sửa số',
    'Cohort retention và chất lượng cohort (H3)',
    'Nền khách ổn định 2020–22 (H6, BTN5)',
    'Bước gãy 2019 là cú sốc thời kỳ (B9, B11)',
], size=13, color=INK, space=8, bullet=True)
tb(s, 6.83, y + 0.02, 5.75, 0.32, 'PHẢI VIẾT LẠI', size=12, color=BAD, bold=True, space=0)
rect(s, 6.83, y + 0.38, 5.75, 2.72, RGBColor(0xFD, 0xEC, 0xEC), rounded=True, adj=0.04)
tb(s, 7.08, y + 0.56, 5.25, 2.4, [
    'H8 — `delivery_days` không phân biệt cả theo vùng',
    'H9 — `returns` không trùng `reviews` đơn nào',
    'Mục 5.1 — `signup_date` có thể là trường ngẫu nhiên',
    'Ba phép kiểm hằng đúng ở Mục 12.2',
    'Ngưỡng K1, K7 và chỉ số K5 ở Mục 8',
], size=13, color=INK, space=8, bullet=True)
callout(s, 0.75, y + 3.3, 11.83, 0.8,
        'Quy tắc rút ra: kết quả **rỗng** trên một trường nghi ngẫu nhiên nói về **bộ sinh dữ liệu**, '
        'không nói về doanh nghiệp — và **không được** chuyển thành khuyến nghị thực hành.',
        color=ACC, size=14)

# ═══════════════════════════════════════════════════════════════════════════════
# 21 · VIỆC CHƯA XONG
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · Giới hạn', 'Việc chưa xong — và vì sao chưa xong',
         'Thiếu công và thiếu dữ liệu là hai chuyện khác nhau, phải nói rõ là chuyện nào')
table(s, 0.75, y + 0.02, 11.83,
      ['', 'Cần gì', 'Làm được không'],
      [['BTN5 · Nền khách đỡ nổi 2023–24?', 'Kiểm định điểm gãy — Chow test / Bai–Perron',
        '✔ Làm được với dữ liệu hiện có'],
       ['BTN6 · Ngân sách nên đi đâu?', 'Bảng chi phí marketing theo kênh và theo kỳ',
        '✖ Bộ dữ liệu KHÔNG CÓ'],
       ['B7 · Dạng hàm của `cohort_year`', 'Thay tuyến tính bằng spline hoặc biến phân loại',
        '✔ Làm được — đã có sẵn script']],
      colw=[4.2, 4.6, 3.03], fs=12.5, rh=0.48,
      align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
      rowcol=[GOOD, BAD, GOOD])
tb(s, 0.75, y + 2.16, 11.83, 0.32, 'BA ĐIỀU THÀNH THẬT NÊN NÓI TRƯỚC KHI BỊ HỎI',
   size=12, color=ACC, bold=True, space=0)
tb(s, 0.75, y + 2.58, 11.83, 1.52, [
    '`promo_first` là **tương quan đã kiểm soát**, không phải nhân quả — matching chỉ cân bằng '
    'biến **đã quan sát**; khách nhạy giá tự chọn vào nhóm khuyến mãi theo đặc điểm không quan sát được.',
    'H6 chỉ dựa trên **n = 3 điểm**. Đủ để nói "chưa thấy dấu hiệu tiếp tục rơi", '
    'chưa đủ để khẳng định chế độ ổn định.',
    '`signup_date` **hỏng nặng** — 73,8% đơn đặt trước ngày đăng ký. Chỉ dùng để bác bỏ khung '
    '"thu nạp hỏng", **không** làm căn cứ định lượng.',
], size=13, hl=DEEP, space=7, bullet=True)

# ═══════════════════════════════════════════════════════════════════════════════
# 22 · NỐI SANG CHƯƠNG MÔ HÌNH
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
y = head(s, 'Phần IV · Nối tiếp', 'Cấu trúc khách hàng nói gì cho chương dự báo?',
         'Ba chế độ, và một ràng buộc mức cho giai đoạn dự báo')
table(s, 0.75, y + 0.02, 11.83,
      ['Giai đoạn', 'Cơ chế bên dưới', 'Hệ quả cho mô hình'],
      [['2014 – 2018', 'Cohort chất lượng cao 2013–2015 (retention ~50%/năm) đang sung sức',
        'Chế độ cao'],
       ['2019', 'Cohort chất lượng cao suy kiệt; cohort thay thế chỉ giữ được ~8%',
        'Điểm gãy'],
       ['2020 – 2022', 'Ổn định ở mức thấp, giữ bởi nền khách lặp lại ~23.000 người',
        'Chế độ thấp, đi ngang']],
      colw=[1.9, 6.7, 3.23], fs=12.5, rh=0.52,
      align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
      rowcol=[GOOD, BAD, ACC])
callout(s, 0.75, y + 2.04, 11.83, 2.06,
        'Ràng buộc mức — **giả định thứ tư** của chương mô hình (kết quả B18):\n'
        'Ngoại suy xu hướng đơn thuần sẽ dự đoán 2023–2024 **tiếp tục giảm**. Nhưng cấu trúc khách hàng '
        'nói **đi ngang quanh mức 2022**.\n'
        'Cụ thể: doanh thu năm **1,16 – 1,23 tỷ** (thang `sales.csv`), nền khách lặp lại '
        'KTC 95% **[22.028 – 23.422]** người.\n'
        'Bổ sung cho ba quyết định kỹ thuật đã có: sample weighting · Fourier seasonality · '
        'không dự báo COGS qua tỷ số cố định.',
        color=ACC, size=13.5)

# ═══════════════════════════════════════════════════════════════════════════════
# 23 · KẾT
# ═══════════════════════════════════════════════════════════════════════════════
s = new()
rect(s, 0, 0, W, H, DEEP)
rect(s, 0, 0, 0.22, H, ACC)
tb(s, 1.25, 1.28, 10.8, 0.35, 'KHI HỘI ĐỒNG HỎI "CHỨNG MINH ĐI"', size=12.5,
   color=RGBColor(0x8F, 0xB8, 0xE0), bold=True, space=0)
tb(s, 1.25, 1.86, 10.7, 2.1,
   ['“Em không chứng minh từng số riêng lẻ.',
    'Em thiết kế để các số **buộc phải khớp nhau** — nếu một số sai thì phép kiểm chéo sẽ vỡ.',
    'Và khi em tự kiểm 31 mục, **tám chỗ đã vỡ thật**. Em sửa cả tám, và đây là chúng.”'],
   size=21, color=WHITE, hl=RGBColor(0x8F, 0xC8, 0xF0), space=10, spacing=1.35)
rect(s, 1.25, 4.34, 1.6, 0.035, ACC)
tb(s, 1.25, 4.7, 11, 1.3, [
    'Tài liệu chính  ·  docs/problem-statement-customer-retention.md',
    'Bản tổng hợp  ·  docs/tong-hop-D2.md',
    'Phản biện 31 mục  ·  docs/phan-bien/cau-hoi-phan-bien-can-kiem.md',
    'Demo 39 ô đã chạy  ·  D2-Demo.ipynb',
], size=13.5, color=RGBColor(0x9E, 0xBE, 0xDC), space=5, font=MONO)
_n[0] = 23

# ── Ghi ra ────────────────────────────────────────────────────────────────────
OUT = 'D2-Slides.pptx'
prs.save(OUT)
print(f'Da tao {OUT} — {len(prs.slides.__iter__.__self__._sldIdLst)} slide')
